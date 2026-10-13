import sys, json, collections, datetime
sys.path.insert(0, '.')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from groups_excluded import GROUPS

OUT = sys.argv[1] if len(sys.argv) > 1 else 'Zocket_TAM_Expansion.xlsx'
rows = json.load(open('merged.json'))
cleanup = json.load(open('tam_cleanup.json')) if __import__('os').path.exists('tam_cleanup.json') else None
Q = [r for r in rows if r['final_status'] == 'Qualified']
R = [r for r in rows if r['final_status'] == 'Rejected']
IND_ORDER = ['Retail, Electronics & Beauty', 'Food, FMCG & QSR', 'Healthcare', 'Health & Wellness', 'BFSI',
             'Home improvement', 'Real Estate', 'Automotive', 'Hospitality', 'Airlines & Aviation', 'Media, Entertainment & Consumer Internet', 'Education']
TAB = {'Retail, Electronics & Beauty': 'Retail, Electronics & Beauty', 'Food, FMCG & QSR': 'Food, FMCG & QSR',
       'Healthcare': 'Healthcare', 'Health & Wellness': 'Health & Wellness', 'BFSI': 'BFSI',
       'Home improvement': 'Home improvement', 'Real Estate': 'Real Estate', 'Automotive': 'Automotive',
       'Hospitality': 'Hospitality', 'Airlines & Aviation': 'Airlines & Aviation', 'Media, Entertainment & Consumer Internet': 'Media, Telecom & Internet', 'Education': 'Education'}
def tier(s): return 'A' if s >= 8 else 'B' if s >= 6 else 'C'
HDR_FILL = PatternFill('solid', fgColor='1F2937'); HDR_FONT = Font(bold=True, color='FFFFFF')

def sheet(wb, title, header, data, widths=None, links=()):
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for c in ws[1]: c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in data: ws.append(row)
    ws.freeze_panes = 'C2' if len(header) > 6 else 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(h) or min(45, max(10, len(str(h)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    for h in links:
        if h in header:
            col = header.index(h) + 1
            for r in range(2, ws.max_row + 1):
                c = ws.cell(r, col)
                if c.value and str(c.value).startswith('http'):
                    c.hyperlink = str(c.value); c.font = Font(color='1D4ED8', underline='single')
    return ws

MAIN = ['Company Name', 'Company Name for Emails', 'Industry', 'Website', 'Company Linkedin Url', 'Company City',
        'Company State', 'Company Country', 'Annual Revenue',
        'Main Industry', 'Revenue (₹ Cr)', 'Revenue FY', 'Revenue Basis', 'Revenue Source', 'Ownership',
        'Parent Group', 'Key Brands', 'Why B2C', 'Consumer Signals', 'Data Confidence',
        'Check: Revenue ≥ ₹500 Cr', 'Check: Consumer-facing', 'Check: India HQ / India-led', 'Check: Not in TAM/Apollo/Exclusions/Groups',
        'ICP Verdict', 'Zocket Fit Score (1-10)', 'Priority Tier', 'Lead Zocket Agent', 'Expansion Agent', 'Source']
W = {'Company Name': 34, 'Company Name for Emails': 28, 'Industry': 30, 'Website': 30, 'Company Linkedin Url': 38,
     'Revenue Source': 40, 'Key Brands': 34, 'Why B2C': 45, 'Consumer Signals': 45, 'Parent Group': 24,
     'Main Industry': 26, 'Lead Zocket Agent': 40, 'Expansion Agent': 30, 'Reason': 60}

def mainrow(r):
    ok = '✓'
    return [r.get('company_name'), r.get('brand_name') or r.get('company_name'), r.get('sub_industry'),
            r.get('website'), r.get('linkedin_url'), r.get('hq_city'), r.get('hq_state'), 'India', r.get('revenue_usd'),
            r.get('industry'), round(r['revenue_inr_cr'], 1), r.get('revenue_fy'), r.get('revenue_basis'),
            r.get('revenue_source_url'), r.get('ownership'), r.get('parent_group'), r.get('key_brands'),
            r.get('b2c_rationale'), r.get('consumer_signals'), r.get('confidence'), ok, ok, ok, ok,
            'Qualified', r['fit_score'], tier(r['fit_score']), r['lead_module'], r['expansion_module'], r.get('source_tag')]

def key(r): return (IND_ORDER.index(r['industry']) if r['industry'] in IND_ORDER else 99, r.get('sub_industry') or '', -r['fit_score'], -(r['revenue_inr_cr'] or 0))
Q.sort(key=key)

wb = Workbook(); ws = wb.active; ws.title = 'README'
today = datetime.date.today().isoformat()
by_ind = collections.Counter(r['industry'] for r in Q)
readme = [
 ['Zocket BrandOS — TAM Expansion (net-new ICP accounts)'], [f'Built {today}'], [],
 ['What this is', f'{len(Q)} net-new companies that meet the ICP and are NOT in the existing 723 TAM, the Apollo export, customers, pipeline, account list or active connections — and do not belong to any excluded business group.'],
 ['ICP rules applied to every row', '1) Revenue ≥ ₹500 Cr, latest FY (FY26 if published, else FY25), with a source link per row. 2) Consumer-facing brand (B2B only if it has a large B2C business; no manufacturers-only). 3) India HQ, or MNC Indian arm only if marketing decisions are made in India (tagged "MNC – India-led"). 4) Not already known; group rule applied both ways (parent, subsidiary, sister, JV, franchisee).'],
 ['One company = one row', 'Each company appears in exactly one industry and one sub-industry (best fit by largest revenue line). Brands owned by the same company are listed in "Key Brands", not as separate rows.'],
 ['Columns A–I', 'Same layout as the existing TAM tabs (Annual Revenue in USD, converted at ₹85/USD).'],
 ['Zocket Fit Score', 'Revenue scale + vertical priority (from Zocket Industry Plays) + consumer/marketing signals + data confidence; MNC arms −1. Priority Tier: A = 8–10, B = 6–7, C = 1–5.'],
 ['Lead Zocket Agent', 'Which agent to lead with in that vertical (per the internal Industry Plays / Positioning pages).'],
 ['Data Confidence', 'High = audited/filed figure; Medium = credible press/aggregator; Low = older or estimated figure.'],
 ['Tabs', 'Summary · All New ICP (combined view) · one tab per industry · Rejected (audit trail) · Excluded Groups · Existing TAM Cleanup'],
 [], ['Net-new qualified by industry'], *[[i, by_ind.get(i, 0)] for i in IND_ORDER], ['TOTAL', len(Q)],
]
for r in readme: ws.append(r)
ws['A1'].font = Font(bold=True, size=14); ws.column_dimensions['A'].width = 32; ws.column_dimensions['B'].width = 140
for row in ws.iter_rows(min_row=4):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')

# Summary
summ = []
for i in IND_ORDER:
    subs = collections.Counter(r.get('sub_industry') for r in Q if r['industry'] == i)
    for s, n in sorted(subs.items(), key=lambda x: -x[1]):
        t = collections.Counter(tier(r['fit_score']) for r in Q if r['industry'] == i and r.get('sub_industry') == s)
        summ.append([i, s, n, t.get('A', 0), t.get('B', 0), t.get('C', 0)])
    summ.append([i, 'Industry total', by_ind.get(i, 0), None, None, None])
summ.append(['ALL', 'GRAND TOTAL', len(Q), sum(1 for r in Q if tier(r['fit_score']) == 'A'), sum(1 for r in Q if tier(r['fit_score']) == 'B'), sum(1 for r in Q if tier(r['fit_score']) == 'C')])
s = sheet(wb, 'Summary', ['Main Industry', 'Sub-industry', 'Net-new companies', 'Tier A', 'Tier B', 'Tier C'], summ, {'Main Industry': 30, 'Sub-industry': 50})
for row in s.iter_rows(min_row=2):
    if row[1].value in ('Industry total', 'GRAND TOTAL'):
        for c in row: c.font = Font(bold=True)

sheet(wb, 'All New ICP (combined)', MAIN, [mainrow(r) for r in Q], W, links=('Website', 'Company Linkedin Url', 'Revenue Source'))
for i in IND_ORDER:
    data = [mainrow(r) for r in Q if r['industry'] == i]
    if data: sheet(wb, TAB[i], MAIN, data, W, links=('Website', 'Company Linkedin Url', 'Revenue Source'))

R.sort(key=lambda r: (r.get('industry') or '', r.get('company_name') or ''))
sheet(wb, 'Rejected (audit)', ['Company Name', 'Main Industry', 'Sub-industry', 'Website', 'Revenue (₹ Cr)', 'Revenue FY', 'Parent Group', 'Reason', 'Revenue Source'],
      [[r.get('company_name'), r.get('industry'), r.get('sub_industry'), r.get('website'), r.get('revenue_inr_cr'), r.get('revenue_fy'), r.get('parent_group'), r.get('final_reason'), r.get('revenue_source_url')] for r in R],
      W, links=('Website', 'Revenue Source'))
sheet(wb, 'Excluded Groups', ['Business group (fully excluded)', 'Why (entity found in customers / pipeline / account list / connections)'],
      [[g, t] for g, t, p in GROUPS if p != '$^'], {'Business group (fully excluded)': 45, 'Why (entity found in customers / pipeline / account list / connections)': 90})
if cleanup:
    sheet(wb, 'Existing TAM Cleanup', cleanup['header'], cleanup['rows'], {h: 40 for h in cleanup['header']})
wb.save(OUT)
print('saved', OUT, len(Q), 'qualified', len(R), 'rejected')

# ---- Extra tabs: Probable (unverified) and Near-miss
import re as _re, importlib as _il
from openpyxl import load_workbook as _lw
import merge as _m
from govt import is_govt as _gv
_wb = _lw(OUT)
_Qn = {_m.nm(r['company_name']) for r in Q}
def _clean(r):
    if _m.nm(r.get('company_name')) in _Qn: return False
    if _gv(r) or _m.group_match(r): return False
    if _re.search(r'already|group|duplicate|b2b|not consumer|government|foreign|prescription|liquidat|defunct|closed|insolv|trust|ngo|subsidiary|holding|obviously|skipped|small', r.get('final_reason',''), _re.I): return False
    k, s = _m.seen_match(r)
    return k != 'strong'
_seen = set(); P = []; N = []
for r in R:
    key = _m.nm(r.get('company_name'))
    if key in _seen or not _clean(r): continue
    rv = r.get('revenue_inr_cr')
    if rv and 300 <= rv < 500: N.append(r); _seen.add(key)
    elif _re.search(r'unverif|not found|no fy|needs verification|not verifiable|mca|old|fy2[0-3]', r.get('final_reason',''), _re.I) and (not rv or rv >= 500):
        P.append(r); _seen.add(key)
def _tab(title, rows, note):
    ws = _wb.create_sheet(title)
    ws.append([note]); ws['A1'].font = Font(bold=True, color='B91C1C')
    hdr = ['Company Name', 'Main Industry', 'Sub-industry', 'Website', 'Parent Group', 'Last known revenue (₹ Cr)', 'Year of figure', 'Why not in main list', 'Source (if any)']
    ws.append(hdr)
    for c in ws[2]: c.fill = HDR_FILL; c.font = HDR_FONT
    for r in sorted(rows, key=lambda r: (r.get('industry') or '', r.get('company_name') or '')):
        ws.append([r.get('company_name'), r.get('industry'), r.get('sub_industry'), r.get('website'), r.get('parent_group'), (r.get('revenue_inr_cr') or None), r.get('revenue_fy'), r.get('final_reason'), r.get('revenue_source_url')])
    for i, w in enumerate([38, 26, 30, 30, 26, 14, 10, 60, 40], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'B3'; ws.auto_filter.ref = f"A2:I{ws.max_row}"
_tab('Probable (unverified rev)', P, f'{len(P)} companies that fit the ICP on profile but have NO public FY24+ revenue figure (private cos that do not publish). Not counted in the main list — verify on MCA/Tofler before outreach.')
_tab('Near-miss ₹300-499 Cr', N, f'{len(N)} ICP-fit companies with verified revenue between ₹300 and ₹499 Cr — just under the bar. Not counted in the main list.')
for _t in ('Probable (unverified rev)', 'Near-miss ₹300-499 Cr'):
    _cur = _wb.sheetnames.index(_t); _tgt = _wb.sheetnames.index('Rejected (audit)')
    _wb.move_sheet(_t, offset=_tgt - _cur)
_wb.save(OUT)
print('extra tabs: probable', len(P), 'near-miss', len(N))

# ---- Mall & Airport brand coverage tab
import glob as _g, json as _j
from load import nm as _nm
_wb = _lw(OUT)
_M = {_nm(r['company_name']): r for r in rows}
_cg = {_nm(r['company_name']): r for r in _j.load(open('cand_gap.json'))}
_ms = []
for _f in sorted(_g.glob('triage/gap_mall_*.jsonl')):
    for _l in open(_f):
        try: _r = _j.loads(_l)
        except Exception: continue
        _k = _nm(_r['company_name']); _m = _M.get(_k); _c = _cg.get(_k, {})
        if not _m:
            _p = _k.replace(' ', '')[:7]
            _m = next((v for kk, v in _M.items() if kk.replace(' ', '').startswith(_p)), None) if len(_p) >= 6 else None
        st = _c.get('status', '')
        if _m and _m['final_status'] == 'Qualified' and not st.startswith(('known','group')): s, d = 'NEW – in this list', f"₹{_m['revenue_inr_cr']:,.0f} Cr {_m.get('revenue_fy')}"
        elif st.startswith('known'):
            src = st.split(':')[1] if ':' in st else ''
            s = 'Already in your TAM (723)' if src == 'TAM723' else 'Excluded (customer/pipeline/account list)'; d = st.split(':', 2)[-1]
        elif st.startswith('group'): s, d = 'Excluded group', st.split(':', 1)[1]
        elif st == 'round1' or st.startswith('in_'):
            s, d = 'Covered by another row', ''
        elif _m: s, d = ('Checked – no public revenue figure (see Probable tab)' if __import__('re').search('unverif|not found|no fy', _m['final_reason'], __import__('re').I) else 'Checked – not added'), _m['final_reason']
        else: s, d = 'Not checked (est. under ₹500 Cr)', ''
        _ms.append([_r.get('brand_name') or _r['company_name'], _r.get('key_brands'), _r['company_name'], _r.get('industry'), _r.get('sub_industry'), s, d])
_ws = _wb.create_sheet('Mall & Airport Brands')
_ws.append(['Brand (as seen in malls/airports)', 'All brands run by operator', 'India operating company', 'Main Industry', 'Sub-industry', 'Status', 'Detail'])
for c in _ws[1]: c.fill = HDR_FILL; c.font = HDR_FONT
for x in sorted(_ms, key=lambda x: (x[5], x[3] or '', x[0] or '')): _ws.append(x)
for i, w in enumerate([30, 50, 40, 26, 30, 34, 60], 1): _ws.column_dimensions[get_column_letter(i)].width = w
_ws.freeze_panes = 'A2'; _ws.auto_filter.ref = _ws.dimensions
_cur = _wb.sheetnames.index('Mall & Airport Brands'); _tgt = _wb.sheetnames.index('Probable (unverified rev)')
_wb.move_sheet('Mall & Airport Brands', offset=_tgt - _cur)
_wb.save(OUT)
import collections as _co
print('mall tab', len(_ms), _co.Counter(x[5] for x in _ms))

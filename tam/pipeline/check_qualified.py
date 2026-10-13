import csv, json, re, sys
sys.path.insert(0, '.')
from load import nm
import merge
from rapidfuzz import fuzz
TAM = list(csv.DictReader(open('inputs/tam_all.csv', encoding='utf-8-sig')))
def keys(n):
    n = re.sub(r'\(.*?\)', ' ', n or '').replace('ABC - ', '')
    k = nm(n); return {x for x in (k, merge.core(k)) if len(x) >= 4}
TK = {}
for t in TAM:
    for k in keys(t['Company Name']) | keys(t['Company Name For Emails']): TK.setdefault(k, t['Company Name'])
def in_tam(*names):
    for n in names:
        for k in keys(n):
            if k in TK: return TK[k]
    for n in names:
        x = nm(n)
        if len(x) < 5: continue
        for k, v in TK.items():
            if len(k) >= 5 and (fuzz.token_sort_ratio(x, k) >= 92): return v
    return None
O = [r for r in json.load(open('tensai/triage.json')) if r['tq'] == 'Qualified']
for r in O:
    hit = in_tam(r['name'])
    print((('TAM: ' + hit) if hit else r['status'])[:60].ljust(60), '|', r['name'][:48])

# ---- final verdicts ----
MAN = {
 'Aero Club Woodland': ('Duplicate', 'Woodland is in the TAM list'),
 'Aisshpra gems And Jewels': ('CAN ADD', '₹1,646 Cr FY25 (CARE) – UP jewellery retail chain'),
 'Ample': ('Already known', 'Ample Technologies is in the account list'),
 'Arundhati Jewellers Pvt. Ltd.': ('CAN ADD', '₹803 Cr FY24 (Acuité/ICRA) – Odisha jewellery retail chain'),
 'BIRLA SUNLIFE INSURANCE': ('Group excluded', 'Aditya Birla group (ABLBL etc. in TAM)'),
 'Darling retail': ('CAN ADD', '₹974 Cr FY23 (ICRA) – TN electronics/furniture retail, 185+ stores'),
 'DHFL Pramerica Life Insurance': ('Duplicate', 'Pramerica Life Insurance is in the TAM list'),
 'Fedbank Financial Services Ltd - Fedfina': ('Duplicate', 'Fedbank Financial Services is in the TAM list'),
 'Flair pens Ltd': ('Duplicate', 'Flair Writing Industries is in the TAM list'),
 'ICICIDirect': ('Group excluded', 'ICICI group (ICICI Home Finance etc. in account list)'),
 'IDBI Federal Life Insurance': ('Duplicate', 'Now Ageas Federal Life – in the TAM list'),
 'Indel Money': ('CAN ADD', '₹606 Cr FY26 total income (Infomerics/NCD prospectus) – gold-loan NBFC'),
 'Indiabulls': ('Duplicate', 'Now Sammaan Capital – in the TAM list'),
 'Janalakshmi Financial Services': ('Already known', 'Now Jana Small Finance Bank – in the account list'),
 'JG Hosiery Pvt. Ltd.- Lifestyle Retail Division': ('CAN ADD', '₹1,241 Cr FY25 (ICRA) – Amul Macho, Sporto innerwear/apparel'),
 'Kotak811': ('Group excluded', 'Kotak Mahindra group'),
 'LazyPay': ('Group excluded', 'PayU group – PayU (Payubiz) is in the account list'),
 'Liberty Videocon General Insurance': ('Duplicate', 'Liberty General Insurance is in the TAM list'),
 'Louis Philippe': ('Group excluded', 'Brand of Aditya Birla Lifestyle Brands – in the TAM list'),
 'Manapuram Finance': ('Already known', 'Manappuram Finance is in the account list'),
 'MetLife': ('Already known', 'PNB MetLife is in the account list'),
 'MR DIY India': ('Not confirmed', 'Only an unofficial estimate (~₹833 Cr) – no filed figure'),
 'Muthoottu Mini Financiers Pvt Ltd': ('CAN ADD', '₹815 Cr FY25 (CARE) – gold-loan NBFC, 948 branches (separate from Muthoot Finance/Fincorp)'),
 'myG': ('Duplicate', 'MYG India is in the TAM list'),
 'N. Ranga Rao & Sons Pvt. Ltd.': ('Duplicate', 'Cycle Pure Agarbathies is in the TAM list'),
 'Officially Ethos': ('Duplicate', 'Ethos Watch Boutiques is in the TAM list'),
 'Marwadi Shares and Finance Ltd.': ('CAN ADD', '₹891 Cr FY25 total operating income (Infomerics, Mar 2026) – Rajkot retail broker, 4 lakh+ clients'),
 'Canon, Inc.': ('Not confirmed', 'Foreign parent; no recent India revenue filing, ~half B2B'),
 'Go Digit General Insurance': ('Duplicate', 'Digit Insurance is in the TAM list'),
 'Hinduja Housing Finance': ('Group excluded', 'Hinduja group (Hinduja Leyland Finance in TAM)'),
 'Opella': ('Duplicate', 'Sanofi Consumer Healthcare India is in the TAM list'),
 'Godrej Housing Finance': ('Group excluded', 'Godrej group'),
 'Piramal Consumer Healthcare': ('Already known', 'In the account list'),
 'ABC - Aditya Birla Housing Finance Limited': ('Group excluded', 'Aditya Birla group'),
 'Aditya Birla Finance Limited': ('Group excluded', 'Aditya Birla group'),
 'Aditya Birla Retail Limited': ('Group excluded', 'Aditya Birla group (now More Retail – also in TAM/exclusions)'),
 'Bajaj Electronics / Electronics Mart': ('Duplicate', 'Electronics Mart India is in the TAM list'),
}
import collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
rows = []
for r in O:
    n = r['name']
    if n in MAN: v, why = MAN[n]
    else:
        hit = in_tam(n); st = r['status']
        if hit: v, why = 'Duplicate', f'In the TAM list as "{hit}"'
        elif st.startswith('Known'):
            src = st.split(': ')[1]; v = 'Already known'
            why = {'ACCOUNT_LIST': 'In the account list', 'PIPELINE': 'In the pipeline', 'CLIENT_META': 'Existing client', 'ACTIVE_CONN': 'Active connection'}.get(src, src)
        elif st == 'Group excluded': v, why = 'Group excluded', r.get('match', '') + ' group'
        else: v, why = '?', st
    rows.append((n, r['trev'], v, why))
c = collections.Counter(x[2] for x in rows); print(c, len(rows))
for x in rows:
    if x[2] in ('?',): print('UNRESOLVED', x)
wb = Workbook(); ws = wb.active; ws.title = 'Qualified 81 – check'
ws.append(['Company (Tensai qualified)', 'Tensai revenue (₹ Cr)', 'Result', 'Reason'])
for cc in ws[1]: cc.font = Font(bold=True, color='FFFFFF'); cc.fill = PatternFill('solid', fgColor='1F2937')
order = ['CAN ADD', 'Not confirmed', 'Duplicate', 'Already known', 'Group excluded']
for x in sorted(rows, key=lambda x: (order.index(x[2]) if x[2] in order else 9, x[0].lower())):
    ws.append(list(x))
    if x[2] == 'CAN ADD':
        for cc in ws[ws.max_row]: cc.fill = PatternFill('solid', fgColor='DCFCE7')
for col, w in zip('ABCD', (46, 28, 16, 90)): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
wb.save('/home/user/Nav/tam/Tensai_Qualified_vs_TAM.xlsx')

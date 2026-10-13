"""Build the re-cleaned TAM: every row of the combined list with a status, reason and revenue evidence."""
import csv, json, glob, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
T = list(csv.DictReader(open('inputs/tam_all.csv', encoding='utf-8-sig')))
E = {r['Company Name'] for r in csv.DictReader(open('inputs/expansion_484.csv', encoding='utf-8-sig'))}
R = json.load(open('reclean/decisions.json'))
EV = {x['row']['Company Name']: x['ev'] for x in json.load(open('reclean/step1.json'))}
RES = {}
for f in glob.glob('reclean/out_*.jsonl'):
    for l in open(f):
        try: x = json.loads(l); RES[x['company']] = x
        except Exception: pass
EXP_REMOVE = {
    'Airtel Payments Bank Ltd': 'Group already with Zocket – Bharti (Bharti AXA in lists)',
    'Bharti Airtel Ltd': 'Group already with Zocket – Bharti (Bharti AXA in lists)',
    'Micro Labs Ltd': 'Not B2C – prescription (Rx) pharma',
    'Jotun India Pvt Ltd': 'Not B2C – industrial/marine paints (removed earlier)',
    'Durian Industries Ltd': 'Below ₹500 Cr revenue (₹100–500 Cr band)',
    'Damro Furniture Pvt Ltd': 'Below ₹500 Cr revenue (₹100–500 Cr band)',
}
STATUS = {'KEEP': 'Keep – ICP', 'REMOVE': 'Remove – not ICP', 'DUP': 'Remove – duplicate', 'ZOCKET': 'Remove – already with Zocket (don\'t cold call)', 'DOUBT': 'Doubtful – check'}
rows = []
for t in T:
    n = t['Company Name']
    rev = src = ''
    if n in E:
        if 'Dealer' in t['Industry']: st, why = 'REMOVE', 'Car dealer – brand advertising run by the carmaker (your rule)'
        elif n in EXP_REMOVE: st, why = ('ZOCKET' if 'Zocket' in EXP_REMOVE[n] else 'REMOVE'), EXP_REMOVE[n]
        else: st, why = 'KEEP', 'Verified in expansion research (revenue ₹500 Cr+ with source, B2C)'
        part = 'Expansion (verified)'
    else:
        part = 'Original Apollo TAM'
        if n in R: st, why = R[n]
        else: st, why = 'KEEP', 'Consumer brand, ₹500 Cr+'
        ev = EV.get(n)
        if ev and ev[1]:
            rev, src = ev[1], ev[3]
            if st == 'KEEP': why = f'Consumer brand – ₹{float(ev[1]):,.0f} Cr ({ev[2]})'
        elif st == 'KEEP':
            why = 'Well-known consumer brand – ₹500 Cr+ (scale widely reported)'
        if st == 'DOUBT' and n in RES:
            x = RES[n]
            st = 'KEEP' if x['verdict'] == 'KEEP' else 'REMOVE'
            rev, src = x.get('revenue_inr_cr') or '', x.get('source_url') or ''
            why = ('Verified – ' if st == 'KEEP' else '') + (x.get('reason') or '')
        OV = {'Kamdhenu Limited': ('REMOVE', 'Not B2C – TMT steel (paints business demerged)'),
              'Del Monte Foods Private Limited': ('DUP', 'Duplicate – owned by Sundrop Brands (Agro Tech Foods) already in TAM'),
              'Sahyadri Hospitals': ('ZOCKET', 'Group already with Zocket – Manipal (now owner) in lists'),
              'Augmont | Gold For All': ('REMOVE', 'Not B2C – revenue mostly B2B bullion trading'),
              'BGS Gleneagles': ('DUP', 'Duplicate – IHH group, same as Fortis already in TAM')}
        if n in OV: st, why = OV[n]
    rows.append(dict(t=t, part=part, st=st, why=why, rev=rev, src=src))
H = list(T[0].keys())
H = [h for h in H if h]
wb = Workbook(); s = wb.active; s.title = 'Summary'
ws = wb.create_sheet('Clean TAM (keep only)')
al = wb.create_sheet('All rows with status')
for w, hdr in ((ws, H + ['Revenue (₹ Cr)', 'Revenue source', 'Part of list']), (al, H + ['Status', 'Reason', 'Revenue (₹ Cr)', 'Revenue source', 'Part of list'])):
    w.append(hdr)
    for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F2937')
COL = {'KEEP': 'DCFCE7', 'REMOVE': 'FEE2E2', 'DUP': 'FEE2E2', 'ZOCKET': 'FEF3C7', 'DOUBT': 'E0E7FF'}
order = ['KEEP', 'DOUBT', 'ZOCKET', 'REMOVE', 'DUP']
for r in sorted(rows, key=lambda r: (order.index(r['st']), r['t']['Industry'], r['t']['Company Name'].lower())):
    base = [r['t'][h] for h in H]
    al.append(base + [STATUS[r['st']], r['why'], r['rev'], r['src'], r['part']])
    al.cell(al.max_row, len(H) + 1).fill = PatternFill('solid', fgColor=COL[r['st']])
    if r['st'] == 'KEEP': ws.append(base + [r['rev'], r['src'], r['part']])
for w in (ws, al):
    for i, wd in enumerate([40, 30, 28, 30, 36, 14, 14, 9, 16, 22, 60, 12, 40, 20][:w.max_column], 1):
        w.column_dimensions[chr(64 + i) if i <= 26 else 'A'].width = wd
    w.freeze_panes = 'B2'; w.auto_filter.ref = w.dimensions
c = collections.Counter((r['part'], r['st']) for r in rows)
rc = collections.Counter((r['st'], r['why'].split(' – ')[0] if r['st'] in ('REMOVE', 'ZOCKET') else '') for r in rows)
s.append(['Re-cleaned TAM – summary']); s['A1'].font = Font(bold=True, size=14)
s.append([f'Rows checked: {len(rows)}']); s.append([])
s.append(['', 'Original Apollo TAM', 'Expansion (verified)', 'Total']); [setattr(x, 'font', Font(bold=True)) for x in s[s.max_row]]
for st in order:
    a, b = c[('Original Apollo TAM', st)], c[('Expansion (verified)', st)]
    s.append([STATUS[st], a, b, a + b])
s.append(['Total', sum(v for (p, _), v in c.items() if p.startswith('Original')), sum(v for (p, _), v in c.items() if p.startswith('Expansion')), len(rows)])
s.append([]); s.append(['Removal reasons', 'Rows']); [setattr(x, 'font', Font(bold=True)) for x in s[s.max_row]]
import re
def bucket(r):
    w = r['why']
    if r['st'] == 'DUP': return 'Duplicate of another row'
    if r['st'] == 'ZOCKET': return 'Already with Zocket (account list / pipeline / group company)'
    if w.startswith('Car dealer'): return 'Car dealer (carmaker runs advertising)'
    if w.startswith('Government'): return 'Government / PSU'
    if w.startswith('Excluded by rule'): return 'Excluded by rule (Zepto, Zomato/Blinkit, malls)'
    if re.search(r'(?i)no verifiable', w): return 'No verifiable revenue figure'
    if re.search(r'(?i)not b2c|b2b|manufactur|rx|foreign|industrial|insolvent|distressed|dsa|hni|msme', w): return 'Not B2C (B2B, manufacturer, exporter, Rx pharma, industrial)'
    return 'Revenue below ₹500 Cr'
why = collections.Counter(bucket(r) for r in rows if r['st'] in ('REMOVE', 'DUP', 'ZOCKET'))
for k, v in why.most_common(): s.append([k, v])
s.column_dimensions['A'].width = 80; s.column_dimensions['B'].width = 20; s.column_dimensions['C'].width = 20
wb.save('/home/user/Nav/tam/Zocket_TAM_Recleaned.xlsx')
print(c); print('KEEP total', sum(1 for r in rows if r['st'] == 'KEEP'))

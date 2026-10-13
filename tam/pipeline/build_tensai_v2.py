import json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
O = json.load(open('tensai/review_out.json'))
Q = {'Qualified – new (adding to Zocket TAM)': 'New ICP account – Zocket will take this',
     'Qualified – already in Zocket TAM': 'ICP fit – already in Zocket data (shared earlier)',
     'Excluded – already in Zocket account list': 'ICP fit – already in Zocket target account list',
     'Excluded – in Zocket pipeline / in talks': 'ICP fit – already in talks / in Zocket pipeline',
     'Excluded – existing Zocket client': 'ICP fit – existing Zocket client (deal closed)'}
RJ_GROUP = 'Group company of an existing Zocket account (parent / subsidiary already covered)'
rows = []
for o in O:
    if o['verdict'] in Q: st, why = 'Qualified', Q[o['verdict']]
    elif o['verdict'].startswith('Excluded – group'): st, why = 'Rejected', RJ_GROUP
    else: st, why = 'Rejected', o['reason']
    det = o['detail'] if (st == 'Rejected' or o['verdict'].startswith('Qualified – new')) else ''
    if why == RJ_GROUP and det and not det.lower().startswith(('same', 'excluded')): det = 'Group: ' + det.replace('Excluded group: ', '')
    rows.append(dict(name=o['name'], vert=o['vert'], parent=o['parent'], st=st, why=why, rev=o['rev'], det=det))
QO = list(Q.values())
RO = ['Below ₹500 Cr revenue', 'Not B2C (B2B, manufacturer, exporter, fintech infra, MSME lender, Rx pharma)',
      'Foreign company – marketing not run from India', 'Serves only HNI / institutional clients', RJ_GROUP,
      'Government / PSU / non-profit', 'Defunct, merged or insolvent', 'Revenue not verifiable (no public figure)', 'Name too generic to identify', 'Low fit for Zocket – limited consumer marketing need']
assert all(r['why'] in QO + RO for r in rows), {r['why'] for r in rows} - set(QO + RO)
N = len(rows); qc = collections.Counter(r['why'] for r in rows if r['st'] == 'Qualified'); rc = collections.Counter(r['why'] for r in rows if r['st'] == 'Rejected')
nq, nr = sum(qc.values()), sum(rc.values())
HDR = PatternFill('solid', fgColor='1F2937'); GREEN = PatternFill('solid', fgColor='DCFCE7'); RED = PatternFill('solid', fgColor='FEE2E2')
wb = Workbook(); s = wb.active; s.title = 'Summary'
s.append(['Tensai lookalike list – Zocket ICP review']); s['A1'].font = Font(bold=True, size=14)
s.append([f'Accounts reviewed: {N}']); s.append(['ICP: Indian consumer-facing (B2C) brand, ₹500 Cr+ revenue (latest FY, filed source), marketing decided in India; not government/non-profit.']); s.append([])
s.append(['Overall', 'Accounts', '% of list']); [setattr(c, 'font', Font(bold=True)) for c in s[s.max_row]]
s.append(['Qualified (ICP fit)', nq, f'{nq/N:.0%}']); s.append(['Rejected (not ICP / not workable)', nr, f'{nr/N:.0%}'])
for c in s[s.max_row - 1]: c.fill = GREEN
for c in s[s.max_row]: c.fill = RED
s.append([])
s.append(['Qualified – breakdown', 'Accounts', '% of list']); [setattr(c, 'font', Font(bold=True)) for c in s[s.max_row]]
for k in QO: s.append([k, qc[k], f'{qc[k]/N:.0%}'])
s.append(['Total qualified', nq, f'{nq/N:.0%}']); s[s.max_row][0].font = Font(bold=True)
s.append([])
s.append(['Rejected – reason', 'Accounts', '% of list']); [setattr(c, 'font', Font(bold=True)) for c in s[s.max_row]]
for k in RO: s.append([k, rc[k], f'{rc[k]/N:.0%}'])
s.append(['Total rejected', nr, f'{nr/N:.0%}']); s[s.max_row][0].font = Font(bold=True)
s.column_dimensions['A'].width = 78; s.column_dimensions['B'].width = 11; s.column_dimensions['C'].width = 11
def sheet(title, sel):
    w = wb.create_sheet(title); w.append(['Account', 'Industry vertical', 'Parent industry', 'Status', 'Reason / remark', 'Est. revenue (₹ Cr)', 'Detail'])
    for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR
    order = QO + RO
    for r in sorted(sel, key=lambda r: (r['st'] != 'Qualified', order.index(r['why']), r['name'].lower())):
        w.append([r['name'], r['vert'], r['parent'], r['st'], r['why'], r['rev'], r['det']])
        w[w.max_row][3].fill = GREEN if r['st'] == 'Qualified' else RED
    for col, wd in zip('ABCDEFG', (40, 26, 16, 11, 62, 20, 80)): w.column_dimensions[col].width = wd
    w.freeze_panes = 'B2'; w.auto_filter.ref = w.dimensions
sheet(f'All {N} accounts', rows)
sheet(f'Qualified ({nq})', [r for r in rows if r['st'] == 'Qualified'])
sheet(f'Rejected ({nr})', [r for r in rows if r['st'] == 'Rejected'])
wb.save('/home/user/Nav/tam/Tensai_Lookalike_Review_Zocket.xlsx')
print(N, nq, nr); print(qc); print(rc)

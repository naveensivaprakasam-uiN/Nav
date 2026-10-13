import csv, json, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
T = list(csv.DictReader(open('inputs/tensai.csv', encoding='utf-8-sig')))
R = json.load(open('tensai/final_rows.json'))
assert len(R) == len(T)
V_NEW_OLD = 'Qualified – added to Zocket TAM'
V_NEW = 'Qualified – new (adding to Zocket TAM)'
BUCKET = [  # (summary group, matcher on reason)
 ('Below ₹500 Cr revenue', lambda r, d: r.startswith('Below ₹500') or r == 'Revenue below ₹500 Cr' or (r.startswith('Revenue below ₹500 Cr or') and not re.search(r'(?i)no |not verif|unverif|conflict', d))),
 ('Revenue not verifiable (no public figure)', lambda r, d: r.startswith('Revenue not verifiable') or r.startswith('Revenue below ₹500 Cr or')),
 ('Not B2C (B2B, manufacturer, exporter, fintech infra, MSME lender, Rx pharma)', lambda r, d: r.startswith('Not B2C')),
 ('Serves only HNI / institutional clients', lambda r, d: r.startswith('Not mass-market')),
 ('Foreign company – marketing not run from India', lambda r, d: r.startswith('Foreign')),
 ('Government / PSU / non-profit', lambda r, d: r.startswith('Government')),
 ('Defunct, merged or insolvent', lambda r, d: r.startswith('Defunct')),
 ('Name too generic to identify', lambda r, d: r.startswith('Name too generic')),
 ('Low fit for Zocket – limited consumer marketing need', lambda r, d: r.startswith('Low fit')),
]
def clean_note(d):
    s = re.split(r'(?<=[.;–])\s+', d or '')
    s = [x for x in s if not re.search(r'(?i)\brule\b|your call|confirm fit|decide if|check to confirm|flagged per|likely fit|icp rules|zocket', x)]
    return ' '.join(s).strip(' –;')
out = []
for t, (n, v, why, d) in zip(T, R):
    if v == V_NEW_OLD: v = V_NEW
    reason = why
    if v == 'Not qualified':
        for g, f in BUCKET:
            if f(why, d or ''): reason = g; break
    out.append(dict(name=n, vert=t['Industry vertical'], parent=t['Parent industry'], verdict=v, reason=reason,
                    rev=t['Est. Annual Revenue (₹ Cr)'], detail=clean_note(d)))
wb = Workbook(); ws = wb.active; ws.title = 'Summary'
H = ['Account', 'Industry vertical', 'Parent industry', 'Zocket verdict', 'Reason', 'Est. revenue (₹ Cr)', 'Detail']
rv = wb.create_sheet('All 1,816 accounts'); rv.append(H)
ORDER = [V_NEW, 'Qualified – already in Zocket TAM', 'Excluded – existing Zocket client', 'Excluded – in Zocket pipeline / in talks',
         'Excluded – already in Zocket account list', 'Excluded – group company already with Zocket', 'Not qualified']
for o in sorted(out, key=lambda o: (ORDER.index(o['verdict']), o['reason'], o['name'].lower())):
    rv.append([o['name'], o['vert'], o['parent'], o['verdict'], o['reason'], o['rev'], o['detail']])
    if o['verdict'] == V_NEW:
        for c in rv[rv.max_row]: c.fill = PatternFill('solid', fgColor='DCFCE7')
for ws_ in (rv,):
    for c in ws_[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F2937')
    for col, w in zip('ABCDEFG', (40, 26, 16, 38, 52, 22, 90)): ws_.column_dimensions[col].width = w
    ws_.freeze_panes = 'B2'; ws_.auto_filter.ref = ws_.dimensions
q = wb.create_sheet('Qualified – new'); q.append(['Account', 'Industry vertical', 'Why it qualifies (revenue, source, business)'])
for o in out:
    if o['verdict'] == V_NEW: q.append([o['name'], o['vert'], o['detail'] or o['reason']])
for c in q[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='166534')
q.column_dimensions['A'].width = 42; q.column_dimensions['B'].width = 30; q.column_dimensions['C'].width = 100
# summary
vc = collections.Counter(o['verdict'] for o in out)
rc = collections.Counter(o['reason'] for o in out if o['verdict'] == 'Not qualified')
ws.append(['Tensai lookalike list – Zocket review']); ws['A1'].font = Font(bold=True, size=14)
ws.append([f'Total accounts reviewed: {len(out)}']); ws.append([])
ws.append(['Outcome', 'Accounts', '% of list'])
for c in ws[ws.max_row]: c.font = Font(bold=True)
for v in ORDER: ws.append([v, vc[v], f"{vc[v]/len(out):.0%}"])
ws.append([]); ws.append(['Why "Not qualified"', 'Accounts', '% of list'])
for c in ws[ws.max_row]: c.font = Font(bold=True)
for g, _ in BUCKET: ws.append([g, rc[g], f"{rc[g]/len(out):.0%}"])
ws.append([]); ws.append(['ICP used: Indian consumer-facing brand, ₹500 Cr+ revenue (latest FY, filed source), India-led marketing; not government/non-profit; not already a Zocket client, prospect or TAM account (incl. parent/subsidiary).'])
ws.column_dimensions['A'].width = 70; ws.column_dimensions['B'].width = 12; ws.column_dimensions['C'].width = 12
wb.move_sheet('Summary', offset=-2)
wb.save('/home/user/Nav/tam/Tensai_Lookalike_Review_Zocket.xlsx')
print(vc); print(rc); print(sum(rc.values()))
json.dump(out, open('tensai/review_out.json', 'w'))
